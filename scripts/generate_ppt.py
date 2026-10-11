import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_master_14_slide_ppt():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Color Palette
    BG_DARK = RGBColor(15, 23, 42)        # Slate 900
    BG_CARD = RGBColor(30, 41, 59)        # Slate 800
    TEXT_LIGHT = RGBColor(248, 250, 252)  # Slate 50
    TEXT_MUTED = RGBColor(148, 163, 184) # Slate 400
    ACCENT_ORANGE = RGBColor(249, 115, 22) # Orange 500
    ACCENT_GREEN = RGBColor(16, 185, 129)  # Emerald 500
    ACCENT_BLUE = RGBColor(59, 130, 246)   # Blue 500
    ACCENT_PURPLE = RGBColor(168, 85, 247) # Purple 500
    ACCENT_RED = RGBColor(239, 68, 68)     # Red 500

    def add_bg(slide, color=BG_DARK):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = color
        bg.line.fill.background()

    def add_header(slide, number_str, title_text, category_text="SITEFORGE VERIFIED AI BUILDER"):
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(10), Inches(0.4))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = f"{number_str} | {category_text}"
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = ACCENT_ORANGE

        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.5), Inches(0.8))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(25)
        p_title.font.bold = True
        p_title.font.color.rgb = TEXT_LIGHT

    slides_data = [
        # SLIDE 1
        {
            "is_title": True,
            "title": "SiteForge",
            "subtitle": "Autonomous Verified AI Website Builder & Hostinger-Style Template Engine",
            "tagline": "Build websites you can trust — verified down to the byte.",
            "notes": "Good morning judges! Today we present SiteForge: an autonomous verified AI website builder that transforms user prompts into production-grade multi-page websites verified by Playwright browser testing and grounded web research."
        },
        # SLIDE 2
        {
            "title": "The Problem: Flaws of Current AI Website Builders",
            "notes": "Existing AI website builders suffer from 3 critical flaws: hallucinating fake facts, generating dead links and 404 pages, and destroying adjacent layout sections during visual edits. SiteForge solves all three.",
            "cards": [
                {
                    "title": "1. Hallucinated Facts",
                    "desc": "Generic AI tools invent dummy phone numbers, non-existent addresses, and unbacked corporate awards.",
                    "tag": "PROBLEM: UNVERIFIED DATA",
                    "color": ACCENT_RED
                },
                {
                    "title": "2. Dead Links & 404 Pages",
                    "desc": "Generated websites frequently contain unlinked CTA buttons, missing subpages, and broken contact forms.",
                    "tag": "PROBLEM: BROKEN CODE",
                    "color": ACCENT_RED
                },
                {
                    "title": "3. Overwrite Damage",
                    "desc": "Prompting an AI to edit one section often mutates or destroys adjacent, perfectly styled sections.",
                    "tag": "PROBLEM: SIDE-EFFECTS",
                    "color": ACCENT_RED
                }
            ]
        },
        # SLIDE 3
        {
            "title": "Our Solution: SiteForge Core Value Proposition",
            "notes": "SiteForge introduces verified web architecture. We combine real-time web research, Playwright automated browser verification, SHA-256 section proofs, and 1-click zero-lockin ZIP export.",
            "cards": [
                {
                    "title": "Grounded RAG Grounding",
                    "desc": "Fetches live web facts via Tavily API. All factual claims carry source citations or an explicit 'Needs Review' marker.",
                    "tag": "ACCURACY GUARANTEE",
                    "color": ACCENT_GREEN
                },
                {
                    "title": "Playwright Link Audit",
                    "desc": "Headless browser crawls every subpage, button, and form across desktop & mobile viewports (0 dead links).",
                    "tag": "VERIFIED CODE",
                    "color": ACCENT_GREEN
                },
                {
                    "title": "Zero Lock-In Export",
                    "desc": "Download standard offline-ready HTML5, CSS3, JS, and JSON spec files with zero vendor lock-in.",
                    "tag": "100% CODE OWNERSHIP",
                    "color": ACCENT_GREEN
                }
            ]
        },
        # SLIDE 4
        {
            "title": "System Architecture & Multi-Agent Pipeline",
            "notes": "Here is our 4-step pipeline: Intake & Web Research via Tavily/DuckDuckGo, Zod Spec synthesis via Groq Llama-3.3-70B, multi-page HTML compilation, and SSE streaming progress logs.",
            "cards": [
                {
                    "title": "Step 01: Web Research Engine",
                    "desc": "Tavily API & DuckDuckGo scrape live facts and business details to construct factual context.",
                    "tag": "RESEARCH PHASE",
                    "color": ACCENT_BLUE
                },
                {
                    "title": "Step 02: Zod Spec Synthesis",
                    "desc": "Groq Llama-3.3-70B synthesizes structured JSON contract validated against Zod schema rules.",
                    "tag": "SPEC CONTRACT",
                    "color": ACCENT_BLUE
                },
                {
                    "title": "Step 03: Code Generation",
                    "desc": "Multi-page compiler renders index.html, about.html, services.html, contact.html, styles.css, script.js.",
                    "tag": "CODE SYNTHESIS",
                    "color": ACCENT_BLUE
                }
            ]
        },
        # SLIDE 5
        {
            "title": "Automated Playwright Link & Responsive Audit (U1)",
            "notes": "Metric U1: Before export is permitted, Playwright headless Chromium crawls every page, button, and form across mobile 375px and desktop 1280px viewports to assert 0 dead links.",
            "cards": [
                {
                    "title": "Headless Browser Crawl",
                    "desc": "Executes automated Playwright instances checking all <a> tags, CTA buttons, and form submit handlers.",
                    "tag": "LINK CRAWLER",
                    "color": ACCENT_ORANGE
                },
                {
                    "title": "Multi-Viewport Audit",
                    "desc": "Simulates rendering at Mobile (375px), Tablet (768px), and Desktop (1280px) breakpoints.",
                    "tag": "RESPONSIVE AUDIT",
                    "color": ACCENT_ORANGE
                },
                {
                    "title": "Human Approval Gate",
                    "desc": "Disables export button until verification checklist passes 100% of link and viewport tests.",
                    "tag": "APPROVAL GATE",
                    "color": ACCENT_ORANGE
                }
            ]
        },
        # SLIDE 6
        {
            "title": "Evidence Grounding & Hallucination Guard (U2, U3)",
            "notes": "Metrics U2 & U3: Grounded RAG citations attach source links to factual claims. Any unverified numerical stat or price is flagged with a visible 'Needs Review' badge in preview.",
            "cards": [
                {
                    "title": "Web Search Grounding",
                    "desc": "Attaches verified source URLs to factual business descriptions, pricing, and service offerings.",
                    "tag": "EVIDENCE PASS",
                    "color": ACCENT_PURPLE
                },
                {
                    "title": "Hallucination Guard",
                    "desc": "Scans generated output for unverified numerical statistics or claims and flags them automatically.",
                    "tag": "HALLUCINATION FLAG",
                    "color": ACCENT_PURPLE
                },
                {
                    "title": "Needs Review Marker",
                    "desc": "Renders prominent UI badges on unverified fields to prompt human review before final export.",
                    "tag": "HUMAN-IN-THE-LOOP",
                    "color": ACCENT_PURPLE
                }
            ]
        },
        # SLIDE 7
        {
            "title": "Visual Studio & SHA-256 Section Locking (U4)",
            "notes": "Metric U4: Zero-collateral section editing uses JSON Patch and SHA-256 checksums to modify target sections while mathematically proving non-target sections remain 100% byte-identical.",
            "cards": [
                {
                    "title": "Point-and-Click Studio",
                    "desc": "Edit text inline, swap theme colors, or adjust layout parameters directly on the live canvas.",
                    "tag": "VISUAL EDITOR",
                    "color": ACCENT_GREEN
                },
                {
                    "title": "Section Locking",
                    "desc": "Lock specific sections to prevent AI prompts from altering verified component layouts.",
                    "tag": "ISOLATION CONTROL",
                    "color": ACCENT_GREEN
                },
                {
                    "title": "SHA-256 Hash Auditing",
                    "desc": "Computes SHA-256 hashes before/after edits to mathematically prove zero collateral damage.",
                    "tag": "HASH INTEGRITY",
                    "color": ACCENT_GREEN
                }
            ]
        },
        # SLIDE 8
        {
            "title": "Hostinger-Style Starter Template Engine (12 Categories)",
            "notes": "SiteForge features a Hostinger-style template gallery with 12 distinct categories. Each template can be inspected in a live preview sandbox and customized with 1 click.",
            "cards": [
                {
                    "title": "12 Category Suite",
                    "desc": "Portfolio, SaaS Startup, Restaurant, Clinic, Law Firm, E-commerce, NGO, Real Estate, Education, Photography, Agency, Blog.",
                    "tag": "TEMPLATE GALLERY",
                    "color": ACCENT_BLUE
                },
                {
                    "title": "Interactive Sandbox",
                    "desc": "Inspect fully rendered static templates in sandboxed iframe viewports prior to customization.",
                    "tag": "LIVE PREVIEW",
                    "color": ACCENT_BLUE
                },
                {
                    "title": "1-Click Customization",
                    "desc": "Pre-loads structured category specifications into the AI builder for 10-second site generation.",
                    "tag": "ONE-CLICK DEPLOY",
                    "color": ACCENT_BLUE
                }
            ]
        },
        # SLIDE 9
        {
            "title": "Student Knowledge Library & RAG Document Ingestion",
            "notes": "The Knowledge Library allows students and researchers to upload PDFs and study notes. SiteForge constructs revision sites and study flashcards grounded in exact source pages.",
            "cards": [
                {
                    "title": "PDF & Note Ingestion",
                    "desc": "Upload course notes, research PDFs, or study guides to construct structured knowledge bases.",
                    "tag": "RAG INGESTION",
                    "color": ACCENT_ORANGE
                },
                {
                    "title": "Grounded Study Guides",
                    "desc": "Generates interactive flashcards and revision sites carrying exact page and chunk citations.",
                    "tag": "FLASHCARD GENERATOR",
                    "color": ACCENT_ORANGE
                },
                {
                    "title": "Zero-Hallucination Fallback",
                    "desc": "Outputs 'Not found in uploaded sources' instead of inventing false information.",
                    "tag": "STRICT CITATIONS",
                    "color": ACCENT_ORANGE
                }
            ]
        },
        # SLIDE 10
        {
            "title": "India-First Localization & Voice Input (U14, U15)",
            "notes": "Metrics U14 & U15: Built for India-first localization with native English, Telugu, and Hindi generation, Web Speech API voice input, INR formatting, UPI QR payments, and WhatsApp integration.",
            "cards": [
                {
                    "title": "Multi-Lingual Generation",
                    "desc": "Native site generation & UI toggles for English, Telugu (తెలుగు), and Hindi (हिन्दी).",
                    "tag": "INDIC LOCALIZATION",
                    "color": ACCENT_GREEN
                },
                {
                    "title": "Voice-to-Website Input",
                    "desc": "Web Speech API integration for hands-free prompt entry in Telugu & English.",
                    "tag": "VOICE INPUT",
                    "color": ACCENT_GREEN
                },
                {
                    "title": "UPI & WhatsApp Integration",
                    "desc": "Generates working UPI QR payment links and direct WhatsApp action contact buttons.",
                    "tag": "LOCAL PAYMENTS",
                    "color": ACCENT_GREEN
                }
            ]
        },
        # SLIDE 11
        {
            "title": "SEO Package, WCAG Accessibility & Security Scanner (U7, U17, U18)",
            "notes": "Metrics U7, U17, U18: Includes automated sitemap.xml, Schema.org JSON-LD data, WCAG 2.2 AA contrast auditing, and honest security scanning for CSP meta tags and script injection.",
            "cards": [
                {
                    "title": "SEO Engine (U17)",
                    "desc": "Generates sitemap.xml, robots.txt, and Schema.org JSON-LD structured data tags automatically.",
                    "tag": "SEARCH OPTIMIZATION",
                    "color": ACCENT_PURPLE
                },
                {
                    "title": "Accessibility Audit (U18)",
                    "desc": "Audits text contrast ratios, image alt attributes, and keyboard focus states with auto-fixes.",
                    "tag": "WCAG 2.2 AA",
                    "color": ACCENT_PURPLE
                },
                {
                    "title": "Security Scanner (U7)",
                    "desc": "Scans for inline script injection, external CDN risks, and enforces Content Security Policies.",
                    "tag": "SECURITY AUDIT",
                    "color": ACCENT_PURPLE
                }
            ]
        },
        # SLIDE 12
        {
            "title": "Project Dashboard & Version Rollback (U5, U9, U12)",
            "notes": "Metrics U5, U9, U12: Project Dashboard displays verification pass rate (100%), version history timeline (v1, v2) with 1-click rollback, and real-time form submission inbox.",
            "cards": [
                {
                    "title": "Project Dashboard (U9)",
                    "desc": "Monitors active projects, verification pass rates (100%), section edit counts, and site analytics.",
                    "tag": "INSIGHTS HUB",
                    "color": ACCENT_BLUE
                },
                {
                    "title": "Version Rollback (U5)",
                    "desc": "Maintains immutable version history timeline (v1, v2, v3) with 1-click instant rollback.",
                    "tag": "REVISION CONTROL",
                    "color": ACCENT_BLUE
                },
                {
                    "title": "Form Submissions Inbox (U12)",
                    "desc": "Logs inbound contact form submissions in real-time under the project dashboard.",
                    "tag": "DATA INBOX",
                    "color": ACCENT_BLUE
                }
            ]
        },
        # SLIDE 13
        {
            "title": "Tech Stack & 100% Free-Tier Architecture",
            "notes": "SiteForge is engineered on a 100% free-tier stack using Next.js 15, Groq Llama-3.3-70B for fast inference, Tavily for web search, and Playwright for automated browser testing.",
            "cards": [
                {
                    "title": "Frontend & Server",
                    "desc": "Next.js 15 App Router, React 19, TypeScript, Tailwind CSS, Lucide Icons.",
                    "tag": "CORE FRAMEWORK",
                    "color": ACCENT_ORANGE
                },
                {
                    "title": "AI & Web Search Stack",
                    "desc": "Groq Free Tier (Llama-3.3-70B), OpenRouter Free Tier, Tavily Search API.",
                    "tag": "LLM & RAG ENGINE",
                    "color": ACCENT_ORANGE
                },
                {
                    "title": "Testing & Automation",
                    "desc": "Playwright Chromium, FFmpeg Static, Google Text-to-Speech API.",
                    "tag": "TESTING SUITE",
                    "color": ACCENT_ORANGE
                }
            ]
        },
        # SLIDE 14
        {
            "title": "Conclusion, Offline ZIP Export & Q&A (U10, U19)",
            "notes": "Metrics U10 & U19: SiteForge outputs 100% offline-ready HTML/CSS/JS ZIP packages with zero vendor lock-in and free deployment guides for Vercel, Netlify, and GitHub Pages. Thank you!",
            "cards": [
                {
                    "title": "Zero Lock-In Export (U10)",
                    "desc": "Download standard offline-ready HTML5, CSS, JS, and JSON spec files in a clean ZIP package.",
                    "tag": "CODE OWNERSHIP",
                    "color": ACCENT_GREEN
                },
                {
                    "title": "Free Deploy Kit (U19)",
                    "desc": "Includes 1-click deployment guides for GitHub Pages, Netlify, Vercel, and Cloudflare Pages.",
                    "tag": "FREE HOSTING",
                    "color": ACCENT_GREEN
                },
                {
                    "title": "Thank You & Q&A",
                    "desc": "Thank you for reviewing SiteForge! We invite your questions and welcome you to try our live demo.",
                    "tag": "LIVE DEMO READY",
                    "color": ACCENT_BLUE
                }
            ]
        }
    ]

    for idx, slide_info in enumerate(slides_data):
        slide = prs.slides.add_slide(blank_layout)
        add_bg(slide)

        notes_slide = slide.notes_slide
        notes_slide.notes_text_frame.text = slide_info["notes"]

        if slide_info.get("is_title"):
            title_box = slide.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(11.33), Inches(1.5))
            tf = title_box.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = slide_info["title"]
            p.font.size = Pt(54)
            p.font.bold = True
            p.font.color.rgb = ACCENT_ORANGE

            sub_box = slide.shapes.add_textbox(Inches(1.0), Inches(3.6), Inches(11.33), Inches(1.0))
            tf_sub = sub_box.text_frame
            tf_sub.word_wrap = True
            p_sub = tf_sub.paragraphs[0]
            p_sub.text = slide_info["subtitle"]
            p_sub.font.size = Pt(22)
            p_sub.font.bold = True
            p_sub.font.color.rgb = TEXT_LIGHT

            tag_box = slide.shapes.add_textbox(Inches(1.0), Inches(5.0), Inches(11.33), Inches(1.0))
            tf_tag = tag_box.text_frame
            tf_tag.word_wrap = True
            p_tag = tf_tag.paragraphs[0]
            p_tag.text = f"Master Presentation Pitch Deck (14 Slides)  |  {slide_info['tagline']}"
            p_tag.font.size = Pt(16)
            p_tag.font.color.rgb = ACCENT_GREEN

        else:
            add_header(slide, f"SLIDE {idx + 1:02d}", slide_info["title"])
            cards = slide_info.get("cards", [])
            card_width = Inches(3.64)
            gap = Inches(0.4)
            start_left = Inches(0.8)
            top_pos = Inches(1.8)
            card_height = Inches(5.0)

            for c_idx, card in enumerate(cards):
                left_pos = start_left + c_idx * (card_width + gap)
                
                shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_pos, top_pos, card_width, card_height)
                shape.fill.solid()
                shape.fill.fore_color.rgb = BG_CARD
                shape.line.color.rgb = card.get("color", ACCENT_BLUE)
                shape.line.width = Pt(1.5)

                tb = slide.shapes.add_textbox(left_pos + Inches(0.2), top_pos + Inches(0.2), card_width - Inches(0.4), card_height - Inches(0.4))
                tf = tb.text_frame
                tf.word_wrap = True

                p_tag = tf.paragraphs[0]
                p_tag.text = card["tag"]
                p_tag.font.size = Pt(10)
                p_tag.font.bold = True
                p_tag.font.color.rgb = card.get("color", ACCENT_BLUE)
                p_tag.space_after = Pt(14)

                p_t = tf.add_paragraph()
                p_t.text = card["title"]
                p_t.font.size = Pt(17)
                p_t.font.bold = True
                p_t.font.color.rgb = TEXT_LIGHT
                p_t.space_after = Pt(14)

                p_d = tf.add_paragraph()
                p_d.text = card["desc"]
                p_d.font.size = Pt(12.5)
                p_d.font.color.rgb = TEXT_MUTED
                p_d.space_after = Pt(6)

    output_dir = os.path.join(os.path.dirname(__file__), '..', 'generated')
    if not os.path.exists(output_dir):
        os.makedirs(output_dir, exist_ok=True)
    
    ppt_path = os.path.join(output_dir, 'SiteForge_Master_14_Slide_Presentation.pptx')
    prs.save(ppt_path)
    print(f"[SUCCESS] Master 14-Slide PowerPoint created at: {ppt_path}")

    # Also update public downloads copy
    public_dir = os.path.join(os.path.dirname(__file__), '..', 'public', 'downloads')
    if os.path.exists(public_dir):
        import shutil
        shutil.copy(ppt_path, os.path.join(public_dir, 'SiteForge_Master_Presentation.pptx'))
        print(f"[SUCCESS] Public download copy updated at public/downloads/SiteForge_Master_Presentation.pptx")

if __name__ == '__main__':
    create_master_14_slide_ppt()
